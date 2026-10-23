
import os
import hashlib
import re
from pathlib import Path

from dotenv import load_dotenv
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_google_genai import (
    GoogleGenerativeAIEmbeddings,
    ChatGoogleGenerativeAI,
)
from langchain_chroma import Chroma
from langchain_community.document_loaders import PyPDFLoader
from pdf2image import convert_from_path
import pytesseract
import docx2txt
from pptx import Presentation
from openpyxl import load_workbook


# --------------------------------------------------
# 1. PROJECT SETTINGS
# --------------------------------------------------

load_dotenv()

API_KEY = os.getenv("GOOGLE_API_KEY")

PROJECT_DIR = Path(__file__).parent
DOCUMENTS_DIR = PROJECT_DIR / "sample_documents"
DOCUMENTS_DIR = PROJECT_DIR / "documents"
CACHE_DIR = PROJECT_DIR / "chroma_cache"

TESSERACT_PATH = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

POPPLER_PATH = (
    r"C:\Users\divya\Downloads\Release-26.09.0-0"
    r"\poppler-26.09.0\Library\bin"
)

pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".pptx", ".xlsx",
    ".csv", ".txt", ".md",
}


# --------------------------------------------------
# 2. LOAD PDF FILES, INCLUDING SCANNED PDFS
# --------------------------------------------------

def load_pdf(file_path: Path) -> list[Document]:
    pages = PyPDFLoader(str(file_path)).load()
    result = []

    for index, page in enumerate(pages):
        text = (page.page_content or "").strip()

        if text:
            page.metadata["source"] = file_path.name
            page.metadata["page"] = index
            result.append(page)
            continue

        page_number = index + 1
        print(f"  OCR: page {page_number}/{len(pages)}")

        images = convert_from_path(
            str(file_path),
            dpi=200,
            first_page=page_number,
            last_page=page_number,
            poppler_path=POPPLER_PATH,
        )

        if images:
            ocr_text = pytesseract.image_to_string(
                images[0], lang="eng"
            ).strip()

            if ocr_text:
                result.append(
                    Document(
                        page_content=ocr_text,
                        metadata={
                            "source": file_path.name,
                            "page": index,
                        },
                    )
                )

    return result


# --------------------------------------------------
# 3. LOAD DIFFERENT DOCUMENT FORMATS
# --------------------------------------------------

def load_document(file_path: Path) -> list[Document]:
    extension = file_path.suffix.lower()
    documents = []

    if extension == ".pdf":
        return load_pdf(file_path)

    elif extension == ".docx":
        text = docx2txt.process(str(file_path)).strip()

        if text:
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": file_path.name,
                        "page": 0,
                    },
                )
            )

    elif extension == ".pptx":
        presentation = Presentation(str(file_path))

        for slide_number, slide in enumerate(
            presentation.slides, start=1
        ):
            text_parts = []

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text_parts.append(shape.text.strip())

                if shape.has_table:
                    for row in shape.table.rows:
                        text_parts.append(
                            " | ".join(
                                cell.text for cell in row.cells
                            )
                        )

            text = "\n".join(text_parts).strip()

            if text:
                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": file_path.name,
                            "page": slide_number - 1,
                            "slide": slide_number,
                        },
                    )
                )

    elif extension == ".xlsx":
        workbook = load_workbook(
            str(file_path),
            read_only=True,
            data_only=True,
        )

        try:
            for sheet in workbook.worksheets:
                rows = []

                for row in sheet.iter_rows(values_only=True):
                    values = [
                        str(value)
                        for value in row
                        if value is not None
                    ]

                    if values:
                        rows.append(" | ".join(values))

                text = "\n".join(rows).strip()

                if text:
                    documents.append(
                        Document(
                            page_content=(
                                f"Sheet: {sheet.title}\n{text}"
                            ),
                            metadata={
                                "source": file_path.name,
                                "sheet": sheet.title,
                                "page": 0,
                            },
                        )
                    )

        finally:
            workbook.close()

    elif extension in {".txt", ".md", ".csv"}:
        text = file_path.read_text(
            encoding="utf-8-sig"
        ).strip()

        if text:
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": file_path.name,
                        "page": 0,
                    },
                )
            )

    return documents


# --------------------------------------------------
# 4. FIND AVAILABLE DOCUMENTS
# --------------------------------------------------


def find_files() -> list[Path]:
    folders = [
        PROJECT_DIR / "sample_documents",
        PROJECT_DIR / "documents",
    ]

    found_files = []

    for folder in folders:
        if not folder.exists():
            folder.mkdir(parents=True)
            continue

        for path in folder.rglob("*"):
            if (
                path.is_file()
                and path.suffix.lower() in SUPPORTED_EXTENSIONS
                and not path.name.startswith("~$")
            ):
                found_files.append(path)

    return sorted(found_files)



# --------------------------------------------------
# 5. LET THE USER SELECT FILES
# --------------------------------------------------

def select_files(files: list[Path]) -> list[Path]:
    print("\nAvailable documents:")

    for index, file_path in enumerate(files, start=1):
        relative_path = file_path.relative_to(PROJECT_DIR)
        print(f"{index}. {relative_path}")

    print("\n0. Search all documents")

    while True:
        choice = input(
            "\nSelect a file number, or enter 0 for all: "
        ).strip()

        if choice == "0":
            return files

        try:
            number = int(choice)

            if 1 <= number <= len(files):
                return [files[number - 1]]

        except ValueError:
            pass

        print("Invalid choice. Enter a listed number.")


# --------------------------------------------------
# 6. CREATE OR REUSE A FILE'S VECTOR DATABASE
# --------------------------------------------------

def cache_path_for(file_path: Path) -> Path:
    # A changed file gets a new cache. An unchanged file
    # can reuse its existing embeddings.
    stat = file_path.stat()

    fingerprint = (
        f"{file_path.resolve()}|"
        f"{stat.st_size}|{stat.st_mtime_ns}"
    )

    digest = hashlib.sha256(
        fingerprint.encode("utf-8")
    ).hexdigest()[:16]

    safe_name = re.sub(
        r"[^A-Za-z0-9_-]", "_", file_path.stem
    )[:40]

    return CACHE_DIR / f"{safe_name}_{digest}"


def create_or_load_store(
    file_path: Path,
    embeddings,
):
    cache_path = cache_path_for(file_path)

    if cache_path.exists():
        try:
            store = Chroma(
                persist_directory=str(cache_path),
                embedding_function=embeddings,
            )

            if store._collection.count() > 0:
                print("  Reusing saved embeddings.")
                return store

        except Exception as error:
            print(f"  Cache could not be reused: {error}")
            print("  Rebuilding this file's database.")

    print("  Reading document...")
    documents = load_document(file_path)

    if not documents:
        print("  No readable text found.")
        return None

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=150,
    )

    chunks = splitter.split_documents(documents)

    if not chunks:
        print("  No text chunks created.")
        return None

    print(f"  Created {len(chunks)} text chunks.")
    print("  Creating embeddings for this file...")

    cache_path.mkdir(parents=True, exist_ok=True)

    try:
        store = Chroma.from_documents(
            documents=chunks,
            embedding=embeddings,
            persist_directory=str(cache_path),
        )
    except Exception:
        # Remove only an incomplete cache for this file.
        import shutil
        shutil.rmtree(cache_path, ignore_errors=True)
        raise

    print("  Saved this file's vector database.")
    return store


# --------------------------------------------------
# 7. PREPARE ONLY THE SELECTED DOCUMENTS
# --------------------------------------------------

def prepare_stores(files: list[Path]):
    if not API_KEY:
        print("ERROR: GOOGLE_API_KEY missing. Check .env.")
        return []

    embeddings = GoogleGenerativeAIEmbeddings(
        model="gemini-embedding-001",
        google_api_key=API_KEY,
    )

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    stores = []

    for file_path in files:
        print(f"\nSelected: {file_path.name}")

        try:
            store = create_or_load_store(
                file_path, embeddings
            )

            if store is not None:
                stores.append(store)

        except Exception as error:
            print(f"  Could not process this file: {error}")

    return stores


# --------------------------------------------------
# 8. FORMAT GEMINI ANSWER AS READABLE TEXT
# --------------------------------------------------

def format_answer(content) -> str:
    if isinstance(content, str):
        return content.strip()

    if isinstance(content, list):
        text_parts = []

        for block in content:
            if isinstance(block, dict):
                if block.get("type") == "text":
                    block_text = block.get("text", "")

                    if block_text:
                        text_parts.append(block_text)

            elif isinstance(block, str):
                text_parts.append(block)

        return "\n".join(text_parts).strip()

    return str(content)


# --------------------------------------------------
# 9. FORMAT DOCUMENT SOURCE LOCATIONS
# --------------------------------------------------

def source_location(doc: Document) -> str:
    source = doc.metadata.get("source", "Unknown file")

    if "slide" in doc.metadata:
        return f"{source} — Slide {doc.metadata['slide']}"

    if "sheet" in doc.metadata:
        return f"{source} — Sheet {doc.metadata['sheet']}"

    if source.lower().endswith(".pdf"):
        page_number = int(doc.metadata.get("page", 0)) + 1
        return f"{source} — Page {page_number}"

    return f"{source} — Document content"


# --------------------------------------------------
# 10. ASK QUESTIONS ABOUT THE SELECTED DOCUMENTS
# --------------------------------------------------

def main():
    files = find_files()

    if not files:
        print("No supported documents found in the documents folder.")
        print("Add PDF, DOCX, PPTX, XLSX, CSV, TXT or MD files.")
        return

    selected_files = select_files(files)

    if not selected_files:
        print("No files selected.")
        return

    print(f"\nSelected {len(selected_files)} file(s).")

    stores = prepare_stores(selected_files)

    if not stores:
        print("No documents could be processed.")
        return

    llm = ChatGoogleGenerativeAI(
        model="gemini-flash-lite-latest",
        google_api_key=API_KEY,
        temperature=0,
    )

    print("\nMulti-Format Document Q&A Assistant ready!")
    print("Ask questions about the selected document(s).")
    print("Type 'exit' to stop.\n")

    while True:
        question = input("Your question: ").strip()

        if question.lower() == "exit":
            print("Assistant stopped.")
            break

        if not question:
            continue

        try:
            # Search each selected file's vector database.
            results = []

            for store in stores:
                matches = (
                    store.similarity_search_with_relevance_scores(
                        question, k=4
                    )
                )

                results.extend(matches)

            # Rank results from all selected files together.
            results.sort(
                key=lambda item: item[1],
                reverse=True,
            )

            docs = [
                doc
                for doc, score in results
                if score >= 0.35
            ][:4]

            if not docs:
                print(
                    "\nI could not find relevant information "
                    "in the selected document(s).\n"
                )
                continue

            context_parts = []

            for doc in docs:
                context_parts.append(
                    f"Source: {source_location(doc)}\n"
                    f"{doc.page_content}"
                )

            context = "\n\n---\n\n".join(context_parts)

            prompt = f"""
You are a document question-answering assistant.

Rules:
1. Answer only using the supplied document context.
2. Do not use outside knowledge or invent facts.
3. If the answer is not supported by the context,
   say: "I could not find this information in the
   selected documents."
4. Keep the answer clear and concise.
5. Mention source file and page, slide, or sheet
   when available.
6. Do not claim that a document defines something
   unless the supplied text actually defines it.

Document context:
{context}

Question:
{question}
"""

            response = llm.invoke(prompt)
            answer = format_answer(response.content)

            print("\nAnswer:")
            print(answer or "No text answer was returned.")

            print("\nRetrieved sources:")
            seen_sources = set()

            for doc in docs:
                location = source_location(doc)

                if location not in seen_sources:
                    print(f"- {location}")
                    seen_sources.add(location)

            print()

        except Exception as error:
            print(f"\nError while answering: {error}\n")
            print(
                "Check internet connection, API quota, "
                "model availability and packages.\n"
            )


# --------------------------------------------------
# 11. PROGRAM START
# --------------------------------------------------

if __name__ == "__main__":
    main()
